# -*- coding: utf-8 -*-
"""
外部工具：读取计数器数据，生成 Excel 表格，并为每项计数器生成一幅折线图。

- 数据来源：data/test/metrics_daily.jsonl（每行一个 JSON，含 date 及各计数器当日增量）
- 输出：单个 .xlsx，所有数据与图表放在同一个工作簿的同一张表
- 每项数据一幅折线图，纵向排列、互不重叠、带中文标题

用法：
    python metrics_excel.py [输出路径]

缺少 openpyxl 时会自动尝试安装（清华源 -> 官方源）。
"""
import json
import os
import subprocess
import sys

DATA_FILE = "data/test/metrics_daily.jsonl"
DEFAULT_OUTPUT = "data/test/metrics_daily.xlsx"

# 字段 -> 中文标题（顺序即工作表列顺序，A=日期，其后为各字段）
FIELD_TITLES = [
    ("mem_created", "记忆创建"),
    ("mem_deleted", "记忆删除"),
    ("link_created", "链接创建"),
    ("link_deleted", "链接删除"),
    ("wordweb_created", "词网创建"),
    ("msg_sent", "消息发送"),
    ("self_ref", "自指消息"),
    ("active_attempt", "主动尝试"),
    ("active_success", "主动成功"),
]


def _ensure_openpyxl():
    """确保 openpyxl 可用；缺失时自动安装（清华源 -> 官方源）。"""
    try:
        import openpyxl  # noqa: F401
        return True
    except ImportError:
        pass
    print("[*] 缺少 openpyxl，正在自动安装...")
    for index_url in (
        "https://pypi.tuna.tsinghua.edu.cn/simple",
        "https://pypi.org/simple/",
    ):
        try:
            subprocess.check_call(
                [sys.executable, "-m", "pip", "install", "openpyxl", "-i", index_url, "-q"]
            )
            import openpyxl  # noqa: F401
            print("[OK] openpyxl 安装成功")
            return True
        except Exception as e:
            print(f"[!] 使用 {index_url} 安装失败: {e}")
    print("[错误] openpyxl 安装失败，请手动执行: pip install openpyxl")
    return False


def _load_records():
    """读取 metrics_daily.jsonl，返回记录列表。"""
    if not os.path.exists(DATA_FILE):
        print(f"[错误] 未找到计数器日志文件: {DATA_FILE}")
        return None
    records = []
    with open(DATA_FILE, "r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            try:
                records.append(json.loads(line))
            except json.JSONDecodeError:
                print(f"[警告] 跳过无法解析的行: {line[:60]}")
                continue
    return records


def main():
    if not _ensure_openpyxl():
        return 1
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.utils import get_column_letter

    records = _load_records()
    if records is None:
        return 1
    if not records:
        print("[提示] 日志文件为空，没有可导出的数据")
        return 0

    output = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_OUTPUT
    os.makedirs(os.path.dirname(output) or ".", exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "指标数据"

    # ---------- 1. 数据表（A=日期，其后为各字段列） ----------
    headers = ["日期"] + [title for _, title in FIELD_TITLES]
    ws.append(headers)
    for rec in records:
        row = [rec.get("date", "")]
        for key, _ in FIELD_TITLES:
            row.append(rec.get(key, 0))
        ws.append(row)

    n = len(records)
    data_last_row = 1 + n

    # ---------- 2. 每项计数器生成一幅折线图，纵向排列不重叠 ----------
    chart_start_col = len(headers) + 2          # 数据表右侧留一列空白后放图表
    rows_per_chart = 20                          # 每张图占约 20 行
    for idx, (key, title) in enumerate(FIELD_TITLES):
        col_idx = 2 + idx                        # 该字段所在列（B=2）
        chart = LineChart()
        chart.title = f"{title}（{key}）"
        chart.style = 12
        chart.y_axis.title = title
        chart.x_axis.title = "日期"
        chart.width = 22
        chart.height = 10

        data = Reference(ws, min_col=col_idx, min_row=1, max_row=data_last_row)
        cats = Reference(ws, min_col=1, min_row=2, max_row=data_last_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)

        start_row = 2 + idx * rows_per_chart
        anchor = f"{get_column_letter(chart_start_col)}{start_row}"
        ws.add_chart(chart, anchor)

    wb.save(output)
    print(f"[OK] 已生成: {output}（{len(records)} 天数据，{len(FIELD_TITLES)} 幅折线图）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
